"""Importa integrantes e guias a partir das planilhas de respostas do Google Forms.

Uso:
    python -m scripts.importar_planilhas \
        --integrantes "Informações de contato (respostas).xlsx" \
        --entidades "Formulário sem título (respostas).xlsx" \
        [--dry-run]

Formatos esperados:
- **Integrantes**: uma linha por pessoa, colunas detectadas por palavra-chave
  ("Nome", "E-mail", "Número de telefone", ...).
- **Guias**: uma linha por pessoa e **uma coluna por linha de trabalho**
  ("Preto(a) Velho(a)", "Caboclo (a)", "Exu", "Pomba Gira", ...); cada célula pode
  trazer mais de um guia ("Zé Caveira e Tranca Ruas").

A desduplicação casa os nomes das duas planilhas (que costumam divergir:
apelidos, nome parcial, erro de digitação) e imprime tudo que foi unificado ou
não encontrado, para revisão.
"""

import argparse
import re
import sys
import unicodedata
from difflib import SequenceMatcher

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database import Base, SessionLocal, engine
from app.models import AreaEnum, CargoEnum, Entidade, Pessoa, TipoGiraEnum

PALAVRAS_NOME = ("nome",)
PALAVRAS_EMAIL = ("email", "e-mail")
PALAVRAS_TELEFONE = ("telefone", "whatsapp", "celular")

# Apelidos usados na planilha de guias que não casam por nome com a de contatos.
APELIDOS = {
    "malu": "Maria Luiza Freitas",
    "dani": "Daniela Pacheco",
}

# Cabeçalho da planilha de guias -> linha de trabalho.
COLUNA_PARA_LINHA = (
    ("preto", TipoGiraEnum.PRETO_VELHO),
    ("caboclo", TipoGiraEnum.CABOCLO),
    ("pomba", TipoGiraEnum.POMBA_GIRA),
    ("exu", TipoGiraEnum.EXU),
    ("ere", TipoGiraEnum.CRIANCA),
    ("crianca", TipoGiraEnum.CRIANCA),
    ("mirim", TipoGiraEnum.MIRIM),
    ("baiano", TipoGiraEnum.BAIANO),
    ("malandro", TipoGiraEnum.MALANDRO),
    ("marinheiro", TipoGiraEnum.MARINHEIRO),
    ("cigano", TipoGiraEnum.CIGANO),
    ("boiadeiro", TipoGiraEnum.BOIADEIRO),
    ("orixa", TipoGiraEnum.ORIXA),
)

# Respostas que significam "não sei / não tenho" e não devem virar guia.
NAO_RESPOSTAS = {
    "",
    "-",
    "--",
    "nao",
    "nao sei",
    "nao sei ainda",
    "nao tenho",
    "nenhum",
    "nenhuma",
    "sem nome",
    "ainda sem nome",
    "n sei",
    "nao lembro",
    "?",
    "x",
    "boa pergunta",
    "tudo",
    "acho que e tudo",
}

# Trechos que indicam comentário em vez de nome de guia.
RUIDOS = (
    "sem nome",
    "nao sei",
    "n sei",
    "nao lembro",
    "afinidade",
    "talvez",
    "nao tenho",
    "ainda nao",
    "desculp",
    "comigo",
)

# "Acho que é o Zé Pilintra" -> "Zé Pilintra"
PREFIXOS_INCERTEZA = re.compile(
    r"^\s*(?:acho que|talvez|parece que)?\s*(?:eu\s+)?(?:tenho|e|é|seja)?\s*"
    r"(?:um|uma|o|a)?\s+",
    re.IGNORECASE,
)

SEPARADORES = re.compile(r"\s*(?:/|,| e | E |&|\+|\n)\s*")


def normalizar(texto: str | None) -> str:
    if not texto:
        return ""
    sem_acento = unicodedata.normalize("NFKD", str(texto))
    sem_acento = "".join(c for c in sem_acento if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", sem_acento).strip().lower()


def titulo(texto) -> str:
    return re.sub(r"\s+", " ", str(texto or "")).strip()


def formatar_telefone(valor) -> str | None:
    """Planilhas do Forms trazem telefone como número (16991569593.0)."""
    if valor is None:
        return None
    if isinstance(valor, float):
        valor = f"{int(valor)}"
    digitos = re.sub(r"\D", "", str(valor))
    return digitos or None


def achar_coluna(cabecalho: list[str], palavras: tuple[str, ...]) -> int | None:
    for indice, coluna in enumerate(cabecalho):
        normalizada = normalizar(coluna)
        if any(palavra in normalizada for palavra in palavras):
            return indice
    return None


def ler_planilha(caminho: str) -> tuple[list[str], list[list]]:
    workbook = load_workbook(caminho, data_only=True)
    planilha = workbook[workbook.sheetnames[0]]
    linhas = list(planilha.iter_rows(values_only=True))
    if not linhas:
        return [], []
    cabecalho = [str(c) if c is not None else "" for c in linhas[0]]
    dados = [list(linha) for linha in linhas[1:] if any(v is not None for v in linha)]
    return cabecalho, dados


def linha_da_coluna(cabecalho_coluna: str) -> TipoGiraEnum | None:
    normalizada = normalizar(cabecalho_coluna)
    for chave, tipo in COLUNA_PARA_LINHA:
        if chave in normalizada:
            return tipo
    return None


def extrair_guias(celula) -> list[str]:
    """Quebra a célula em nomes de guias, descartando "não sei" e comentários."""
    if celula is None:
        return []
    texto = str(celula)
    # remove parênteses ("(acho)", "(ele é um sereio ...)") e emojis
    texto = re.sub(r"\([^)]*\)", " ", texto)
    texto = re.sub(r"[^\w\sÀ-ÿ/,+&\-\n]", " ", texto)

    guias = []
    for parte in SEPARADORES.split(texto):
        nome = PREFIXOS_INCERTEZA.sub("", titulo(parte)).strip()
        normalizado = normalizar(nome)
        if not nome or normalizado in NAO_RESPOSTAS:
            continue
        if any(marca in normalizado for marca in RUIDOS):
            continue
        # respostas longas são comentário do formulário, não nome de guia
        if len(normalizado) < 3 or len(nome.split()) > 5:
            continue
        guias.append(nome)
    return guias


class Deduplicador:
    """Casa nomes das planilhas com as pessoas já conhecidas."""

    def __init__(self, pessoas: list[Pessoa]):
        self.pessoas: dict[str, Pessoa] = {normalizar(p.nome): p for p in pessoas}
        self.unificados: list[str] = []
        self.novos: list[str] = []

    def registrar(self, pessoa: Pessoa) -> None:
        self.pessoas[normalizar(pessoa.nome)] = pessoa

    def buscar(self, nome: str) -> Pessoa | None:
        chave = normalizar(APELIDOS.get(normalizar(nome), nome))
        if chave in self.pessoas:
            pessoa = self.pessoas[chave]
            if chave != normalizar(nome):
                self.unificados.append(f'"{nome}" -> "{pessoa.nome}" (apelido)')
            return pessoa

        tokens = set(chave.split())
        candidatos = [
            pessoa
            for existente, pessoa in self.pessoas.items()
            if tokens
            and (tokens <= set(existente.split()) or set(existente.split()) <= tokens)
        ]
        if not candidatos:
            # erro de digitação: mesma quantidade de tokens e alta similaridade
            candidatos = [
                pessoa
                for existente, pessoa in self.pessoas.items()
                if len(existente.split()) == len(chave.split())
                and SequenceMatcher(None, existente, chave).ratio() >= 0.92
            ]

        if len(candidatos) == 1:
            pessoa = candidatos[0]
            if normalizar(pessoa.nome) != normalizar(nome):
                self.unificados.append(f'"{nome}" -> "{pessoa.nome}"')
            return pessoa
        if len(candidatos) > 1:
            nomes = ", ".join(p.nome for p in candidatos)
            self.unificados.append(
                f'AMBÍGUO: "{nome}" casa com vários ({nomes}) — criado separado'
            )
        return None


def importar_integrantes(
    db: Session, caminho: str, dedup: Deduplicador
) -> dict[str, int]:
    cabecalho, linhas = ler_planilha(caminho)
    col_nome = achar_coluna(cabecalho, PALAVRAS_NOME)
    if col_nome is None:
        raise SystemExit(
            f"Não encontrei a coluna de nome em {caminho}. Cabeçalho: {cabecalho}"
        )
    col_email = achar_coluna(cabecalho, PALAVRAS_EMAIL)
    col_telefone = achar_coluna(cabecalho, PALAVRAS_TELEFONE)

    criados = atualizados = ignorados = 0
    for linha in linhas:
        nome = titulo(linha[col_nome]) if col_nome < len(linha) else ""
        if not nome:
            ignorados += 1
            continue
        # planilhas com nome em CAIXA ALTA ficam padronizadas
        if nome.isupper():
            nome = nome.title()

        def valor(indice):
            if indice is None or indice >= len(linha) or linha[indice] is None:
                return None
            return titulo(linha[indice]) or None

        email = valor(col_email)
        telefone = formatar_telefone(
            linha[col_telefone]
            if col_telefone is not None and col_telefone < len(linha)
            else None
        )

        pessoa = dedup.buscar(nome)
        if pessoa:
            pessoa.email = pessoa.email or email
            pessoa.telefone = pessoa.telefone or telefone
            atualizados += 1
            continue

        pessoa = Pessoa(
            nome=nome,
            email=email,
            telefone=telefone,
            cargo=CargoEnum.FIXO,
            area=AreaEnum.MEDIUNIDADE,
            mensalidade=0.0,
            ativo=1,
        )
        db.add(pessoa)
        db.flush()
        dedup.registrar(pessoa)
        criados += 1

    return {"criados": criados, "atualizados": atualizados, "ignorados": ignorados}


def importar_entidades(
    db: Session, caminho: str, dedup: Deduplicador
) -> dict[str, int]:
    cabecalho, linhas = ler_planilha(caminho)
    col_medium = achar_coluna(cabecalho, PALAVRAS_NOME)
    if col_medium is None:
        raise SystemExit(
            f"Não encontrei a coluna de nome em {caminho}. Cabeçalho: {cabecalho}"
        )

    colunas_guias = [
        (indice, linha_da_coluna(coluna))
        for indice, coluna in enumerate(cabecalho)
        if indice != col_medium and linha_da_coluna(coluna) is not None
    ]
    if not colunas_guias:
        raise SystemExit(
            f"Nenhuma coluna de linha de trabalho reconhecida em {cabecalho}"
        )

    criadas = duplicadas = mediuns_novos = linhas_ignoradas = 0
    for linha in linhas:
        nome_medium = titulo(linha[col_medium]) if col_medium < len(linha) else ""
        if not nome_medium:
            linhas_ignoradas += 1
            continue

        pessoa = dedup.buscar(nome_medium)
        if not pessoa:
            pessoa = Pessoa(
                nome=nome_medium,
                cargo=CargoEnum.FIXO,
                area=AreaEnum.MEDIUNIDADE,
                mensalidade=0.0,
                ativo=1,
            )
            db.add(pessoa)
            db.flush()
            dedup.registrar(pessoa)
            dedup.novos.append(f'"{nome_medium}" (só na planilha de guias)')
            mediuns_novos += 1

        for indice, tipo in colunas_guias:
            celula = linha[indice] if indice < len(linha) else None
            for nome_guia in extrair_guias(celula):
                ja_existe = any(
                    normalizar(e.nome) == normalizar(nome_guia) and e.tipo == tipo
                    for e in pessoa.guias
                )
                if ja_existe:
                    duplicadas += 1
                    continue
                # append na coleção (e não db.add) para a checagem acima ver
                # também os guias criados nesta mesma execução
                pessoa.guias.append(Entidade(nome=nome_guia, tipo=tipo))
                criadas += 1

    return {
        "guias_criados": criadas,
        "guias_duplicados_ignorados": duplicadas,
        "mediuns_criados_pela_planilha_de_guias": mediuns_novos,
        "linhas_ignoradas": linhas_ignoradas,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--integrantes", help="Planilha de informações de contato (.xlsx)"
    )
    parser.add_argument("--entidades", help="Planilha de guias (.xlsx)")
    parser.add_argument(
        "--dry-run", action="store_true", help="Não grava nada no banco"
    )
    args = parser.parse_args()

    if not args.integrantes and not args.entidades:
        parser.error("Informe --integrantes e/ou --entidades")

    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        dedup = Deduplicador(db.query(Pessoa).all())

        if args.integrantes:
            print("Integrantes:", importar_integrantes(db, args.integrantes, dedup))
        if args.entidades:
            print("Guias:", importar_entidades(db, args.entidades, dedup))

        if dedup.unificados:
            print("\nNomes unificados (revise):")
            for aviso in dedup.unificados:
                print(f"  - {aviso}")
        if dedup.novos:
            print("\nIntegrantes criados sem ficha de contato:")
            for aviso in dedup.novos:
                print(f"  - {aviso}")

        if args.dry_run:
            db.rollback()
            print("\n[dry-run] nada foi gravado.")
        else:
            db.commit()
            print("\nImportação concluída.")
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    sys.exit(main())
